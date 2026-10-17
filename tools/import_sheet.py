#!/usr/bin/env python3
"""Rebuild the app's schedule data from the "Wer Wann" spreadsheet.

Reads every "Wer Wann <year>" tab of the workbook (2022 onwards) and creates
volunteers, sessions and Yes / No answers in PocketBase. Until the
switch-over the spreadsheet stays the source of truth, so the importer can
wipe the app's schedule data and load it again from the latest sheet.

Usage:
    pip install openpyxl
    PB_PASSWORD='superuser-password' python3 tools/import_sheet.py Einsatzdaten.xlsx \\
        --url http://127.0.0.1:8090 --admin-email you@example.com \\
        [--reset] [--admins NAME,...] [--key-holders NAME,...] [--dry-run]

Without --reset, running it again updates what is there: volunteers are
matched by name, sessions by date, answers by session and volunteer.
With --reset, all volunteers, sessions and answers are deleted first.

How the sheet is read:
- Row 2 holds the column titles. Volunteer columns are the named columns
  between the date (and the 🔑 / 🧑‍🔧 status columns) and the first money or
  shift lead column ("ScheffIn", "Scheff", "Total", "Bargeld").
- "Name (Reserve)" becomes "Name"; the same name in different years is
  the same volunteer.
- Yep becomes yes; (Ev) and Nope become no (the app has no maybe).
- Tabs with a 🔑 column: a session without a key status is closed, and its
  remark becomes the reason.
- "ScheffIn" becomes the session's lead, "Bemerkungen" its notes, and the
  money and visitor columns are kept per session in `sheet_totals` for the
  statistics.
- When tabs overlap (e.g. 1 Jan 2026 in the 2025 tab), the later tab wins.
- Key holders: the names given with --key-holders, otherwise the yellow name
  headers of the latest tab.
- Volunteers named in the latest tab are active, everybody else inactive.
- Email and phone come from the "Team Email" tab when the first name is
  unique there ("Nicht mehr aktiv" makes the volunteer inactive). Others get
  a placeholder address that an admin fixes in the app before they can log in.
"""
import argparse
import datetime as dt
import json
import os
import re
import sys
import urllib.error
import urllib.parse
import urllib.request

import openpyxl

ANSWERS = {"yep": "yes", "(ev)": "no", "nope": "no"}
END_OF_VOLUNTEERS = ("scheff", "total", "bargeld")
STATUS_HEADERS = ("🔑", "🧑")
YELLOW = "FFFFFF00"
TOTALS = {
    # sheet_totals key: how the column title starts (lower case)
    "total": ("total",),
    "cash": ("5.- rote kasse", "bargeld"),
    "twint": ("5-liber twint",),
    "material": ("mat & spenden",),
    # 2022 has "Anzahl Besucher" per session and "Besucher" per quarter.
    "visitors": ("anzahl besucher", "besucher"),
}


def clean_name(name):
    return re.sub(r"\s*\(.*?\)\s*", " ", str(name)).strip()


def slug(name):
    return re.sub(r"[^a-z0-9]+", ".", name.lower()).strip(".")


def number(value):
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def parse_tab(ws):
    """Return the volunteer names (with key colour) and the sessions of one tab."""
    header = {c.column: str(c.value).strip() for c in ws[2] if c.value is not None}
    lower = {col: title.lower() for col, title in header.items()}

    def find(*prefixes, exact=False):
        """First column matching the first prefix that matches any column."""
        for p in prefixes:
            for col, title in sorted(lower.items()):
                if title == p if exact else title.startswith(p):
                    return col
        return None

    key_col = find("🔑")
    lead_col = find("scheff")
    remark_col = find("bemerk")
    end = min(col for col, title in lower.items() if title.startswith(END_OF_VOLUNTEERS))
    volunteers = {
        col: (clean_name(title), ws.cell(2, col).fill.fgColor.rgb == YELLOW)
        for col, title in header.items()
        if 1 < col < end and not title.startswith(STATUS_HEADERS)
    }
    total_cols = {key: find(*prefixes, exact=key == "visitors") for key, prefixes in TOTALS.items()}

    sessions = []
    for row in ws.iter_rows(min_row=3):
        date = row[0].value
        if not isinstance(date, dt.datetime):
            continue
        cell = lambda col: ws.cell(row[0].row, col).value if col else None  # noqa: E731
        remark = str(cell(remark_col) or "").strip()
        closed = key_col is not None and cell(key_col) is None
        answers = {}
        for col, (name, _) in volunteers.items():
            value = str(cell(col) or "").strip().lower()
            if value in ANSWERS:
                answers[name] = ANSWERS[value]
        totals = {key: number(cell(col)) for key, col in total_cols.items() if number(cell(col)) is not None}
        sessions.append({
            "date": date.date(),
            "closed": closed,
            "remark": remark,
            "lead": clean_name(cell(lead_col) or "").lower(),
            "answers": answers,
            "totals": totals,
        })
    return list(volunteers.values()), sessions


def read_contacts(wb):
    """Contacts from the "Team Email" tab, keyed by lower-case first name.

    The schedule only has first names, so a contact is used when its first
    name is unique in the tab.
    """
    if "Team Email" not in wb.sheetnames:
        return {}
    contacts, seen = {}, {}
    for row in wb["Team Email"].iter_rows(min_row=4, max_col=5, values_only=True):
        name, email, _address, phone, remark = (list(row) + [None] * 5)[:5]
        if not name:
            continue
        first = str(name).split()[0].lower()
        seen[first] = seen.get(first, 0) + 1
        contacts[first] = {
            "email": str(email).strip() if email and "@" in str(email) else None,
            "phone": re.sub(r"\s*\(?privat\)?\s*", "", str(phone or ""), flags=re.I).strip(),
            "inactive": "nicht mehr aktiv" in str(remark or "").lower(),
        }
    return {k: v for k, v in contacts.items() if seen[k] == 1}


class Api:
    def __init__(self, url, dry_run):
        self.url = url.rstrip("/")
        self.token = None
        self.dry_run = dry_run
        self.fake_ids = 0

    def call(self, method, path, body=None):
        req = urllib.request.Request(
            self.url + path,
            method=method,
            data=json.dumps(body).encode() if body is not None else None,
            headers={"content-type": "application/json", **({"Authorization": self.token} if self.token else {})},
        )
        try:
            with urllib.request.urlopen(req) as res:
                return json.loads(res.read() or b"null")
        except urllib.error.HTTPError as e:
            raise SystemExit(f"{method} {path} failed: {e.code} {e.read().decode()}")

    def login(self, email, password):
        res = self.call("POST", "/api/collections/_superusers/auth-with-password", {"identity": email, "password": password})
        self.token = res["token"]

    def all(self, collection, fields="id"):
        items, page = [], 1
        while True:
            q = urllib.parse.urlencode({"page": page, "perPage": 500, "fields": fields})
            res = self.call("GET", f"/api/collections/{collection}/records?{q}")
            items += res["items"]
            if page >= res["totalPages"]:
                return items
            page += 1

    def write(self, collection, existing_id, data):
        if self.dry_run:
            self.fake_ids += 1
            return {"id": existing_id or f"dry{self.fake_ids}", **data}
        if existing_id:
            return self.call("PATCH", f"/api/collections/{collection}/records/{existing_id}", data)
        return self.call("POST", f"/api/collections/{collection}/records", data)

    def delete_all(self, collection):
        items = self.all(collection)
        if not self.dry_run:
            for item in items:
                self.call("DELETE", f"/api/collections/{collection}/records/{item['id']}")
        return len(items)


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("xlsx")
    p.add_argument("--url", default="http://127.0.0.1:8090")
    p.add_argument("--admin-email", required=True, help="PocketBase superuser; password in PB_PASSWORD")
    p.add_argument("--reset", action="store_true", help="delete all volunteers, sessions and answers first")
    p.add_argument("--admins", default="", help="comma separated names that get admin rights")
    p.add_argument("--key-holders", default="", help="comma separated names with a key (default: yellow names)")
    p.add_argument("--default-hours", default="19:00-21:00", help="hours of sessions on days without session rule")
    p.add_argument("--dry-run", action="store_true")
    args = p.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    tabs = sorted((n for n in wb.sheetnames if re.fullmatch(r"Wer Wann \d{4}", n)), key=lambda n: n[-4:])
    if not tabs:
        sys.exit("No 'Wer Wann <year>' tabs found.")
    contacts = read_contacts(wb)

    # Collect everything first, later tabs override earlier ones.
    names, sessions = {}, {}
    latest_names, yellow = set(), set()
    for tab in tabs:
        tab_volunteers, tab_sessions = parse_tab(wb[tab])
        for name, is_yellow in tab_volunteers:
            names.setdefault(name.lower(), name)
        for e in tab_sessions:
            sessions[e["date"]] = e
        latest_names = {n.lower() for n, _ in tab_volunteers}
        yellow = {n.lower() for n, is_yellow in tab_volunteers if is_yellow}
        print(f"{tab}: {len(tab_volunteers)} volunteers, {len(tab_sessions)} sessions")

    split = lambda s: {x.strip().lower() for x in s.split(",") if x.strip()}  # noqa: E731
    key_holders = split(args.key_holders) or yellow
    admins = split(args.admins)
    for unknown in (key_holders | admins) - set(names):
        print(f"warning: {unknown!r} is not in any tab")

    api = Api(args.url, args.dry_run)
    api.login(args.admin_email, os.environ.get("PB_PASSWORD") or sys.exit("Set PB_PASSWORD."))

    if args.reset:
        for collection in ("availability", "sessions", "users"):
            print(f"deleted {api.delete_all(collection)} {collection}")

    existing_users = {u["name"].lower(): u["id"] for u in api.all("users", "id,name")}
    existing_sessions = {e["date"]: e["id"] for e in api.all("sessions", "id,date")}
    existing_answers = {(a["session"], a["volunteer"]): a["id"] for a in api.all("availability", "id,session,volunteer")}
    hours = {r["weekday"]: (r["start"], r["end"]) for r in api.all("opening_rules", "weekday,start,end")}
    default_hours = tuple(args.default_hours.split("-"))

    volunteer_ids = {}
    for key, name in sorted(names.items()):
        contact = contacts.get(key, {})
        email = contact.get("email") or f"{slug(name)}@example.invalid"
        data = {
            "name": name,
            "email": email,
            "emailVisibility": True,
            "phone": contact.get("phone", ""),
            "has_key": key in key_holders,
            "is_admin": key in admins,
            "active": key in latest_names and not contact.get("inactive"),
        }
        volunteer_ids[key] = api.write("users", existing_users.get(key), data)["id"]
        flags = [f for f, on in (("key", data["has_key"]), ("admin", data["is_admin"]), ("inactive", not data["active"])) if on]
        print(f"volunteer {name:14} {' '.join(flags):18} {'NO EMAIL, fix in the app' if email.endswith('.invalid') else ''}")

    answers = 0
    for date, e in sorted(sessions.items()):
        start, end = hours.get(date.isoweekday(), default_hours)
        data = {
            "date": date.isoformat(),
            "start": start,
            "end": end,
            "status": "closed" if e["closed"] else "open",
            "closed_reason": (e["remark"] or "geschlossen") if e["closed"] else "",
            "notes": "" if e["closed"] else e["remark"],
            "lead": volunteer_ids.get(e["lead"], ""),
            "sheet_totals": e["totals"] or None,
        }
        session_id = api.write("sessions", existing_sessions.get(date.isoformat()), data)["id"]
        for name, answer in e["answers"].items():
            volunteer_id = volunteer_ids[name.lower()]
            api.write(
                "availability",
                existing_answers.get((session_id, volunteer_id)),
                {"session": session_id, "volunteer": volunteer_id, "answer": answer},
            )
            answers += 1
    print(f"{len(volunteer_ids)} volunteers, {len(sessions)} sessions, {answers} answers{' (dry run)' if args.dry_run else ''}")


if __name__ == "__main__":
    main()
