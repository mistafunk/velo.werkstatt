"""Tests for reading the spreadsheet layouts. Run: python3 -m unittest discover tools"""
import datetime as dt
import unittest

import openpyxl
from openpyxl.styles import PatternFill

from import_sheet import parse_tab

YELLOW = PatternFill("solid", fgColor="FFFFFF00")


def sheet(header, rows, yellow=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["title"])
    ws.append(header)
    for col, title in enumerate(header, start=1):
        if title in yellow:
            ws.cell(2, col).fill = YELLOW
    for row in rows:
        ws.append(row)
    return ws


class ParseTab(unittest.TestCase):
    def test_current_layout(self):
        ws = sheet(
            ["Datum", "🔑", "🧑‍🔧", "Volunteer A", "Volunteer B (Reserve)", "Volunteer C", "ScheffIn", "Total", "5.- rote Kasse",
             "5-liber Twint", "Mat & Spenden", "Besucher", "Besucher pro Q", "Bemerkungen"],
            [
                [dt.datetime(2026, 2, 23), "✅", 2, "Yep", "(Ev)", "Nope", "Volunteer A ", 78, None, 35, 43, 7, None, "Pedale offen"],
                [dt.datetime(2026, 4, 6), None, 0, "Nope", None, None, None, 0, None, None, None, None, None, "Ostermontag"],
                ["Total", None, None, None, None, None, None, 3236],
            ],
            yellow={"Volunteer A", "Volunteer B (Reserve)"},
        )
        volunteers, sessions = parse_tab(ws)
        self.assertEqual(volunteers, [("Volunteer A", True), ("Volunteer B", True), ("Volunteer C", False)])
        self.assertEqual(len(sessions), 2)
        first, holiday = sessions
        self.assertEqual(first["answers"], {"Volunteer A": "yes", "Volunteer B": "maybe", "Volunteer C": "no"})
        self.assertEqual(first["lead"], "volunteer a")
        self.assertEqual(first["totals"], {"total": 78, "twint": 35, "material": 43, "visitors": 7})
        self.assertEqual(first["remark"], "Pedale offen")
        self.assertFalse(first["closed"])
        self.assertTrue(holiday["closed"])
        self.assertEqual(holiday["remark"], "Ostermontag")

    def test_2022_layout_without_status_columns(self):
        ws = sheet(
            ["Datum", "Volunteer A", "Volunteer B", "Bargeld 5.-", "Anzahl Besucher", "Besucher"],
            [[dt.datetime(2022, 3, 28), "Yep", "Nope", 25, 8, 65]],
        )
        volunteers, sessions = parse_tab(ws)
        self.assertEqual([n for n, _ in volunteers], ["Volunteer A", "Volunteer B"])
        self.assertEqual(sessions[0]["answers"], {"Volunteer A": "yes", "Volunteer B": "no"})
        self.assertEqual(sessions[0]["totals"], {"cash": 25, "visitors": 8})
        self.assertFalse(sessions[0]["closed"])


if __name__ == "__main__":
    unittest.main()
